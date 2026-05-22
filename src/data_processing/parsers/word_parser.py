import os
from typing import Any, Dict

try:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
except ImportError:
    Document = None  # type: ignore
    DocxPackageNotFoundError = None  # type: ignore

from .base import BaseParser


def _make_unique_columns(columns: list) -> list:
    """对重复列名添加 _N 后缀以确保唯一性。"""
    seen: dict[str, int] = {}
    result = []
    for col in columns:
        if col in seen:
            cur = seen[col] + 1
            seen[col] = cur
            result.append(f"{col}_{cur}")
        else:
            seen[col] = 0
            result.append(col)
    return result


def _table_to_records(header: list, rows: list) -> list[dict]:
    """将表格头+行数据转为字典列表，自动处理重复列名。"""
    columns = _make_unique_columns(header)
    return [dict(zip(columns, row)) for row in rows]


def _extract_table_from_doc(table) -> list[dict]:
    """从 python-docx Table 对象提取记录列表。"""
    rows_data = []
    for row in table.rows:
        rows_data.append([cell.text.strip() for cell in row.cells])

    if not rows_data:
        return []
    return _table_to_records(rows_data[0], rows_data[1:])


class WordParser(BaseParser):
    def parse(self, file_path: str) -> Dict[str, Any]:
        if Document is None:
            raise ImportError(
                "python-docx is required. Install it with: pip install python-docx"
            )

        try:
            doc = Document(file_path)
        except DocxPackageNotFoundError as e:
            raise FileNotFoundError(str(e))

        # 提取段落文本
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)

        # 提取表格
        all_tables = []
        for table in doc.tables:
            records = _extract_table_from_doc(table)
            if records:
                all_tables.append(records)

        print(f"Parsing Word file: {file_path}")
        return {
            "content": full_text.strip(),
            "metadata": {
                "type": "docx",
                "source": os.path.basename(file_path),
                "paragraph_count": len(paragraphs),
                "table_count": len(doc.tables),
            },
            "tables": all_tables,
        }
