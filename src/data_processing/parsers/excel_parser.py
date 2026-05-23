import os
from typing import Any, Dict

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

from .base import BaseParser


def _make_unique_columns(columns: list) -> list:
    """对重复列名添加 _N 后缀以确保唯一性。"""
    seen: dict[str, int] = {}
    result = []
    for col in columns:
        if col in seen:
            cur = seen[col] + 1
            seen[col] = cur
            result.append(f"{col}_{cur}")
        else:
            seen[col] = 0
            result.append(col)
    return result


def _table_to_records(header: list, rows: list) -> list[dict]:
    """将表格头+行数据转为字典列表，自动处理重复列名。"""
    columns = _make_unique_columns(header)
    return [dict(zip(columns, row)) for row in rows]


def _sheet_to_records(ws) -> list[dict]:
    """将 openpyxl Worksheet 转为记录列表（第一行为列名）。"""
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c) if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return []

    data_rows = []
    for row in rows_iter:
        # 跳过全空行
        if all(cell is None or (isinstance(cell, str) and cell.strip() == "") for cell in row):
            continue
        data_rows.append([str(c) if c is not None else "" for c in row])

    return _table_to_records(header, data_rows)


class ExcelParser(BaseParser):
    def parse(self, file_path: str) -> Dict[str, Any]:
        if openpyxl is None:
            raise ImportError(
                "openpyxl is required. Install it with: pip install openpyxl"
            )

        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"Excel 文件不存在: {file_path}")

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(
                f"无法打开 Excel 文件 (可能已损坏或格式不符): {file_path}\n"
                f"  详情: {e}"
            ) from e

        sheet_text_parts = []
        all_tables = []

        for sheet_name in wb.sheetnames:
            try:
                ws = wb[sheet_name]
                records = _sheet_to_records(ws)
                if records:
                    all_tables.append({"sheet": sheet_name, "data": records})
                    lines = [f"[Sheet: {sheet_name}]"]
                    for row in records[:5]:
                        line = " | ".join(str(v) for v in row.values())
                        lines.append(line)
                    if len(records) > 5:
                        lines.append(f"... ({len(records)} 行)")
                    sheet_text_parts.append("\n".join(lines))
            except Exception as e:
                print(f"  [警告] Sheet '{sheet_name}' 解析异常: {e}")
                continue

        wb.close()

        print(f"  Excel 解析完成: {os.path.basename(file_path)} "
              f"({len(wb.sheetnames)} 个 Sheet, {len(all_tables)} 个表格)")
        return {
            "content": "\n\n".join(sheet_text_parts).strip(),
            "metadata": {
                "type": "excel",
                "source": os.path.basename(file_path),
                "sheet_count": len(wb.sheetnames),
                "sheets": wb.sheetnames,
            },
            "tables": all_tables,
        }
